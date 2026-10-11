"""Extract source-faithful V21 appendices and the supplied official timetable.

Usage: python tools/import-election-documents.py TIMETABLE.docx V21.docx
PDF rendering is performed separately with Chromium from the generated HTML.
"""
from pathlib import Path
from html import escape
import json
import re
import sys
from docx import Document
from docx.text.paragraph import Paragraph
from docx.table import Table

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'public/documents/v21-forms'
OUT.mkdir(parents=True, exist_ok=True)
W = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'

source = Document(sys.argv[2])
blocks = list(source.element.body)
start = next(i for i, b in enumerate(blocks) if b.tag == W+'p' and Paragraph(b, source).text.strip() == 'MPU 2026 Election Forms')
appendix_start = next(i for i, b in enumerate(blocks) if i > start and b.tag == W+'p' and Paragraph(b, source).text.startswith('Appendix A:'))
# Preserve Word's original formatting, relationships, embedded artwork and sections.
booklet = Document(sys.argv[2])
for b in list(booklet.element.body)[:start]:
    booklet.element.body.remove(b)
booklet.save(OUT / 'MPU-2026-V21-Appendices-Only.docx')

def paragraph_html(element):
    p = Paragraph(element, source)
    pieces = []
    for run in p.runs:
        text = escape(run.text).replace('\n', '<br>').replace('\t', ' &nbsp; ')
        if run.bold:
            text = '<strong>'+text+'</strong>'
        if run.italic:
            text = '<em>'+text+'</em>'
        pieces.append(text)
        for blip in run.element.findall('.//{http://schemas.openxmlformats.org/drawingml/2006/main}blip'):
            rid = blip.get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed')
            if rid:
                part = source.part.related_parts[rid]
                name = Path(str(part.partname)).name
                (OUT / name).write_bytes(part.blob)
                pieces.append('<img src="'+escape(name)+'" alt="Source form illustration">')
    if not pieces:
        pieces = [escape(p.text)]
    return '<p>'+(''.join(pieces) or '&nbsp;')+'</p>'

def block_html(element):
    if element.tag == W+'p':
        return paragraph_html(element)
    if element.tag == W+'tbl':
        table = Table(element, source)
        rows = []
        for index, row in enumerate(table.rows):
            cells = []
            # Work from actual XML cells so merged Word cells are not duplicated.
            for cell in row._tr.tc_lst:
                span = cell.tcPr.gridSpan.val if cell.tcPr is not None and cell.tcPr.gridSpan is not None else 1
                shade = cell.find('.//'+W+'shd')
                fill = shade.get(W+'fill') if shade is not None else None
                style = ' style="background:#'+fill+'"' if fill and re.fullmatch('[0-9A-Fa-f]{6}', fill) else ''
                tag = 'th' if index == 0 else 'td'
                content = ''.join(paragraph_html(p) for p in cell if p.tag == W+'p')
                cells.append(f'<{tag} colspan="{span}"{style}>{content}</{tag}>')
            rows.append('<tr>'+''.join(cells)+'</tr>')
        return '<table>'+''.join(rows)+'</table>'
    return ''

groups = []
for b in blocks[appendix_start:]:
    text = Paragraph(b, source).text if b.tag == W+'p' else ''
    if text.startswith('Appendix'):
        groups.append({'title': text, 'blocks': []})
    if groups:
        groups[-1]['blocks'].append(b)

css = '''body{overflow-wrap:anywhere;font:11pt Arial,sans-serif;color:#172b43;margin:0;background:#edf2f8}header,main{max-width:1100px;margin:auto;padding:20px}nav{display:flex;gap:12px;flex-wrap:wrap}a,button{font:inherit}nav a,button{padding:10px 14px;background:#215e92;color:white;border:0;border-radius:5px;text-decoration:none;cursor:pointer}.notice{padding:12px;background:#fff5d7}article{background:white;padding:22px;margin:20px 0;border:1px solid #ccd7e3}h1{font-size:24px}h2{font-size:19px}p{line-height:1.35;white-space:pre-wrap;margin:7px 0}table{border-collapse:collapse;width:100%;table-layout:fixed;margin:12px 0}td,th{border:1px solid #555;padding:6px;vertical-align:top;overflow-wrap:anywhere}td p,th p{margin:0;line-height:1.2}th{background:#e8eff5}img{max-width:100%;max-height:170px;object-fit:contain}article>p:first-of-type{font-weight:bold;font-size:16px} #appendix-5 table{font-size:9pt}#appendix-5 td,#appendix-5 th{padding:4px}article[data-appendix="F"] table,article[data-appendix="G"] table,article[data-appendix="H"] table{font-size:9pt}article[data-appendix="F"] td,article[data-appendix="G"] td,article[data-appendix="H"] td{padding:3px} @media(max-width:600px){header,main{padding:10px}article{padding:10px}table{font-size:8pt}td,th{padding:3px}} @page{size:letter;margin:0.45in} @page wide{size:letter landscape;margin:0.45in} @media print{article[data-appendix="B"],article[data-appendix="I"],#appendix-5{page:wide}body{background:white;color:black;font-size:10pt}header,.notice,.toolbar{display:none!important}main{padding:0;max-width:none}article{border:0;padding:0;margin:0;break-before:page}article:first-child{break-before:auto}tr{break-inside:avoid}img{max-height:140px}body.single article{display:none}body.single article.selected{display:block;break-before:auto}}'''
articles = []
manifest = []
for i, group in enumerate(groups):
    identifier = 'appendix-'+str(i+1)
    letter = re.search(r'Appendix\s*:?\s*([A-L])', group['title']).group(1)
    articles.append(f'<article id="{identifier}" data-appendix="{letter}"><div class="toolbar"><h2>{escape(group["title"])}</h2><button data-print="{identifier}">Print this form</button></div>'+''.join(block_html(b) for b in group['blocks'])+'</article>')
    manifest.append({'id':identifier,'title':group['title']})
links = ''.join(f'<li><a href="#{x["id"]}">{escape(x["title"])}</a></li>' for x in manifest)
html = '<!doctype html><html lang="en"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><base href="/documents/v21-forms/"><title>MPU 2026 V21 Printable Appendix Forms</title><style>'+css+'</style></head><body><header><nav><a href="/">← Portal Home</a><a href="/voter-education/">Voter Education</a><a href="MPU-2026-V21-Appendices-Only.docx" download>Download Word booklet</a><a href="MPU-2026-V21-Appendices-Only.pdf" download>Download PDF booklet</a><button id="print-all">Print all forms</button></nav><h1>MPU 2026 — V21 Appendix Forms</h1><p class="notice">V21 REVIEW DRAFT — templates for review. Sample ballots are illustrations only, not official voting ballots. Source wording is preserved. The main rules section is excluded.</p><ul>'+links+'</ul></header><main>'+''.join(articles)+'</main><script>function clearPrint(){document.body.classList.remove("single");document.querySelectorAll(".selected").forEach(e=>e.classList.remove("selected"))}document.querySelectorAll("[data-print]").forEach(b=>b.addEventListener("click",()=>{clearPrint();document.body.classList.add("single");document.getElementById(b.dataset.print).classList.add("selected");window.print()}));document.getElementById("print-all").addEventListener("click",()=>{clearPrint();window.print()});window.addEventListener("afterprint",clearPrint);</script><script defer src="/assets/section-navigation.js"></script><script defer src="/assets/designer-credit.js"></script></body></html>'
(OUT/'index.html').write_text(html)
(OUT/'manifest.json').write_text(json.dumps(manifest,indent=2))

time_doc = Document(sys.argv[1])
rows = time_doc.tables[1].rows if 'DATE' not in time_doc.tables[0].rows[0].cells[0].text else time_doc.tables[0].rows
schedule = [('2026-09-30',None,'general'),('2026-10-01','2026-10-02','publication'),('2026-10-03','2026-10-09','education'),('2026-10-11',None,'publication'),('2026-10-11','2026-10-17','register-review'),('2026-10-18',None,'deadline'),('2026-10-19','2026-10-25','validation'),('2026-10-26','2026-10-28','screening'),('2026-10-31',None,'publication'),('2026-10-31','2026-11-03','review'),('2026-11-03T23:59:00-05:00',None,'deadline'),('2026-11-04',None,'publication'),('2026-11-05','2026-11-14','campaign'),('2026-11-08',None,'meeting'),('2026-11-15',None,'election')]
assert len(rows)-1 == len(schedule), (len(rows),len(schedule))
events = []
for i,(row,(start_date,end,phase)) in enumerate(zip(rows[1:],schedule)):
    date,title,description = [c.text.strip() for c in row.cells]
    event = dict(id='official-v2-'+str(i+1),title=title,start=start_date,displayDate=date+', 2026',description=description,published=True,phaseType=phase,source='Official Election Timetable FINAL v2')
    if end: event['end']=end
    if phase == 'register-review':event.update(actionLabel='Check My Record',actionUrl='/register-review/')
    if phase == 'education':event.update(actionLabel='Open Voter Education',actionUrl='/voter-education/')
    events.append(event)
target = ROOT/'public/assets/official-election-timetable-v2.json'
target.write_text(json.dumps(events,indent=2))
runtime = ROOT/'public/assets/runtime.js'
text = runtime.read_text()
a = text.index('  const seedEvents=')
b = text.index('\n\n  const defaultSettings=', a)
old_seed = text[a:b].replace('const seedEvents=', 'const previousSeedEvents=',1)
if 'official-v2-1' not in text[a:b]:
    text = text[:a]+old_seed+'\n  const seedEvents='+json.dumps(events,indent=2)+';'+text[b:]
    text = text.replace('function ensure(){if(!localStorage.getItem(EVENT_KEY))', 'function ensure(){if(JSON.stringify(read(EVENT_KEY,null))===JSON.stringify(previousSeedEvents))write(EVENT_KEY,seedEvents);if(!localStorage.getItem(EVENT_KEY))')
    runtime.write_text(text)
print(f'Created {len(groups)} printable appendix sections and {len(events)} timetable events.')
