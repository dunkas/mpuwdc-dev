"""Create the constitutional review question document from the supplied checklist.

Questions are editorial prompts, not new rules or an inferred constitutional answer key.
The five original Day 1 Google Form questions are reproduced separately verbatim.
"""
from pathlib import Path
from html import escape
import json
from docx import Document
from docx.shared import Inches, Pt

ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'public/documents/constitutional-questions'
OUT.mkdir(parents=True,exist_ok=True)
questions=[
('Constitution version / ratification status','Which exact Constitution version is duly ratified and controlling for this election? What record confirms its adoption, amendments, effective date and ratification status?'),
('Candidate eligibility','Who may nominate or be nominated under the controlling Constitution? Which age, membership, attendance, financial and office-specific qualifications must the EC verify?'),
('Active-member status','How does the Constitution define an active member and good standing? Which official records establish dues status and the applicable attendance requirement?'),
('Minimum age','What minimum age must a candidate satisfy, and what evidence establishes the candidate’s age?'),
('Minimum membership duration','What minimum membership duration is required to contest office? How is the membership start date verified?'),
('Candidate attendance requirement','Which attendance requirements apply to candidates with different lengths of membership? Which meetings and official attendance records count?'),
('Candidate presence at election','When does the Constitution require candidate presence? How does V21 document presence at the call for election and address an absent candidate?'),
('Financial standing','Which constitutional financial obligations and timing requirements apply to candidates and active members? Which MPU officer verifies the records, and how does the October 18 administrative records cutoff operate without changing constitutional requirements?'),
('Voter eligibility','Who is constitutionally entitled to vote? How does the EC verify the voter register against official membership, attendance, financial and LGA records?'),
('One vote per office','What voting entitlement does each eligible member have for each office? How are general offices and LGA-specific Board seats distinguished?'),
('Proxy voting','Which constitutional provision governs proxy voting? How do accreditation and voting-table controls ensure that eligible members vote personally?'),
('Board of Directors voting','Which Board positions are voted on by the membership generally, and which are voted on only by the applicable LGA? How are LGA eligibility and ballot isolation verified?'),
('Elective offices','What is the complete constitutional list of elected offices and seats? How does that list correspond to the office sections in V21 and the ballot?'),
('Nomination procedure','Which written declaration or nomination methods does the Constitution authorize? How are the nominator’s and nominee’s eligibility, the specific office and attributable submission verified?'),
('Required nominators','How many eligible nominators does the Constitution require? Is there a valid constitutional basis for requiring more than one?'),
('Nomination fees','Does the Constitution authorize a nomination fee? If a fee is proposed, what controlling provision or properly authorized decision supports it?'),
('Secret ballot','Which provision requires ballot secrecy? How do ticket handling, signatures, ballot issuance, assistance and recordkeeping avoid identifying a voter’s marked selections?'),
('Election method / winning threshold','What winning threshold does the Constitution require? How do V21 procedures calculate valid votes, abstentions and exclusions for each office without changing that threshold?'),
('Election quorum','What quorum applies to an election, and how is the controlling active-member total established? What constitutionally authorized action applies if quorum is not met?'),
('Election timing','When must the election take place under the Constitution? Does the planned November election correspond to the required quarterly meeting and election cycle?'),
('Term of office','What term of office applies to elected officers? What dates define the beginning and end of that term?'),
('Assumption of office','When do newly elected officers assume office? What transition steps must occur between the election and that effective date?'),
('Term limits / reelection','Which consecutive-service and reelection restrictions apply to each office? Which prior-service records establish a candidate’s eligibility?'),
('Multiple offices','Which restrictions apply to holding more than one elected office, and how is committee service distinguished? How does V21 enforce one operative candidacy without amending constitutional office-holding restrictions?'),
('Spousal office restriction','What does the Constitution provide about husband-and-wife simultaneous office holding? What disclosure, screening record and competent decision resolve an actual or potential conflict?'),
('Presidential rotation','What is the constitutional presidential-rotation order among the LGAs? Which LGA is entitled to produce the President for this cycle, and what official record confirms that determination?'),
('Failure of eligible LGA to produce a President','What happens if the LGA whose turn it is does not produce an eligible presidential candidate? Which constitutional conditions govern passing the opportunity to the next LGA?'),
('Election postponement','Under what emergency or exigency conditions may an election be postponed? Which authority, consultation, decision, notice and limits does the Constitution require?'),
('Electoral Committee authority','How must the Electoral Committee be appointed and approved or ratified? Does the official election file contain the required appointment and General Assembly authority records?'),
('Candidate withdrawal','Does the Constitution prescribe a withdrawal process? Which parts of V21 are administrative procedures, and how are verified withdrawals, deadlines, late withdrawals and historical records handled without creating a new constitutional disqualification?'),
('Tie resolution','What constitutional provision or competent-authority decision governs a tie or failure to obtain the required majority? If the matter is a genuine gap, how is Article XII followed rather than assuming authority?'),
('Election complaints / petitions','Which constitutional grievance or petition process applies to each type of election concern? How are administrative record corrections distinguished from misconduct allegations and constitutional-violation petitions?'),
('Appeal authority','Which authority may review the particular decision or grievance? What constitutional basis establishes jurisdiction before the Board of Directors is treated as an appeal body?'),
('Vacancies / special elections','Which constitutional procedures apply to vacancies, temporary appointments and subsequent elections? How are those procedures kept distinct from the regular general-election process?'),
('Campaign conduct','Which existing constitutional member-conduct provisions govern election and campaign conduct? How does V21 apply them neutrally and uniformly?'),
('Campaign sanctions','Which sanctions and disciplinary processes are constitutionally authorized? Which authority may impose them, and what process is required before any candidate disqualification or other substantive consequence?'),
('Transition / handover','What records, property and responsibilities must outgoing officers transfer? Which constitutional steps, responsible persons and timeline govern formal handover?'),
('Unforeseen election matters','When the Constitution does not cover a matter, what Article XII process applies? What roles do the Executive Council, General Assembly and any urgent President-led consultation and reporting process have?'),
]
original=json.loads((ROOT/'public/assets/day1-original-questions.json').read_text())
doc=Document();section=doc.sections[0];section.page_width=Inches(8.5);section.page_height=Inches(11)
section.top_margin=section.bottom_margin=Inches(.65);section.left_margin=section.right_margin=Inches(.7)
doc.styles['Normal'].font.name='Calibri';doc.styles['Normal'].font.size=Pt(11)
doc.add_heading('MPU 2026 — Constitutional Questions',0)
doc.add_paragraph('EC review and member education | Prepared October 11, 2026')
intro='Coverage: all 37 electoral issues in the uploaded “Mpu revised electiion needs” checklist, its additional Constitution-version/ratification question, and all five original questions exposed by the supplied Day 1 Google Form. These are review questions, not adopted rules or an authoritative answer key.'
doc.add_paragraph(intro)
doc.add_paragraph('V21 is the selected election-procedures source. Constitutional answers must be supported by the duly ratified Constitution. The checklist cites “MPU Constitution Review 2023-2”; its assertions should not be treated as proof of adoption. The original V21 source’s draft notice is retained in the portal.')
doc.add_heading('Part I — Questions for Constitutional Verification',1)
html_questions=[]
for number,(topic,question) in enumerate(questions,1):
    doc.add_heading(f'{number}. {topic}',2);doc.add_paragraph(question)
    doc.add_paragraph('Controlling Article / Section and verification record: ______________________________')
    html_questions.append(f'<section class="question"><h3>{number}. {escape(topic)}</h3><p>{escape(question)}</p><p class="reference">Controlling Article / Section and verification record: ______________________________</p></section>')
doc.add_page_break();doc.add_heading('Part II — Original Day 1 Member Questions',1)
doc.add_paragraph('Question wording and choices below are copied from the supplied MPU 2026 KNOW THE RULES — DAY 1: VOTER READY Google Form. The public form does not expose its answer key. No answer has been invented or graded here.')
html_original=[]
for number,q in enumerate(original,len(questions)+1):
    doc.add_heading(f'{number}. {q["text"]}',2)
    for choice in q['options']:doc.add_paragraph('☐ '+choice)
    html_original.append('<section class="question"><h3>'+str(number)+'. '+escape(q['text'])+'</h3>'+''.join('<p>☐ '+escape(c)+'</p>' for c in q['options'])+'</section>')
doc.add_heading('Source and Review Notes',1)
notes=[
 'Uploaded reference: Mpu revised electiion needs.pdf — constitutional compliance checklist, including its final version/ratification-status note.',
 'Selected procedures: MPU-2026-Election-Rules-and-Forms-V21-Updated-Timetable-Draft.docx. Selection for the portal does not change the wording or adoption notice in the source.',
 'Original Day 1 Google Form: https://docs.google.com/forms/d/e/1FAIpQLSdjCWm_eV_--m-wntlw-GWCBjKRglvi2ehPvxrush4CTsdqBw/viewform',
 'The education landing page and Apps Script client were inspected. The complete server-side randomized question bank and answer keys were not exposed by those pages; this document does not claim to reproduce that hidden bank.',
 'EC review should record the controlling provision, constitutional version, evidence, competent authority and decision. Nothing in this question list creates a new eligibility requirement, fee, sanction, appeal jurisdiction or election procedure.',
]
for note in notes:doc.add_paragraph(note)
doc.save(OUT/'MPU-2026-Constitutional-Questions.docx')
css='body{font:11pt Arial,sans-serif;color:#18324e;margin:0;background:#edf2f8}header,main{max-width:950px;margin:auto;padding:20px}main{background:white}nav{display:flex;gap:12px;flex-wrap:wrap}nav a,button{padding:10px;background:#215e92;color:white;border:0;border-radius:5px;text-decoration:none;font:inherit;cursor:pointer}p{line-height:1.45;overflow-wrap:anywhere}h3{font-size:13pt;margin-bottom:6px}.question{break-inside:avoid;margin-bottom:20px}.reference{font-size:9pt;color:#4b5664}.original{break-before:page}@page{size:letter;margin:.65in}@media print{header{display:none}body{background:white;color:black}main{padding:0;max-width:none}h1{font-size:21pt}}'
html='<!doctype html><html lang="en"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>MPU Constitutional Questions</title><style>'+css+'</style></head><body><header><nav><a href="/admin/">← EC Admin Center</a><a href="/voter-education/">Voter Education</a><a href="/documents/constitutional-questions/MPU-2026-Constitutional-Questions.docx" download>Download Word</a><a href="/documents/constitutional-questions/MPU-2026-Constitutional-Questions.pdf" download>Download PDF</a><button onclick="window.print()">Print Questions</button></nav></header><main><h1>MPU 2026 — Constitutional Questions</h1><p>EC review and member education | Prepared October 11, 2026</p><p>'+escape(intro)+'</p><p>V21 is the selected election-procedures source. Answers require verification against the duly ratified Constitution; the uploaded checklist is a reference document.</p><h2>Part I — Questions for Constitutional Verification</h2>'+''.join(html_questions)+'<div class="original"><h2>Part II — Original Day 1 Member Questions</h2><p>Original question wording and choices are preserved. The answer key was not exposed, so no answers or scores are supplied.</p>'+''.join(html_original)+'</div><h2>Source and Review Notes</h2>'+''.join('<p>'+escape(n)+'</p>' for n in notes)+'</main><script defer src="/assets/section-navigation.js"></script></body></html>'
(OUT/'index.html').write_text(html)
print(f'Created {len(questions)} review questions + {len(original)} original Day 1 questions = {len(questions)+len(original)} total.')
